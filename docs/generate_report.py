#!/usr/bin/env python3
"""Генерация отчёта по этапу 2 курсовой работы."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "coursework_stage2_autoservice.docx"


def set_run_font(run, *, size=14, bold=False, name="Times New Roman"):
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = name
    r = run._element
    rPr = r.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:eastAsia"), name)


def add_center(doc, text, *, size=14, bold=False, space_after=6):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold)
    return p


def add_body(doc, text, *, size=14, bold=False, first_line=True):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.5
    if first_line:
        p.paragraph_format.first_line_indent = Cm(1.25)
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold)
    return p


def add_heading_custom(doc, text, level=1):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.5
    run = p.add_run(text)
    set_run_font(run, size=16 if level == 1 else 14, bold=True)
    return p


def add_code(doc, text):
    for line in text.splitlines():
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.left_indent = Cm(0.5)
        run = p.add_run(line if line else " ")
        set_run_font(run, size=10, name="Courier New")


def read_sql(name: str, max_chars: int | None = None) -> str:
    text = (ROOT / "sql" / name).read_text(encoding="utf-8")
    if max_chars and len(text) > max_chars:
        return text[:max_chars] + "\n-- ... (полный код в файле sql/" + name + ")"
    return text


def build():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(3)
    section.right_margin = Cm(1.5)

    # ---------- Титульный лист (как в этапе 1) ----------
    add_center(doc, "МИНОБРНАУКИ РОССИИ", bold=True, size=12, space_after=0)
    add_center(
        doc,
        "ФЕДЕРАЛЬНОЕ ГОСУДАРСТВЕННОЕ АВТОНОМНОЕ",
        bold=True,
        size=12,
        space_after=0,
    )
    add_center(
        doc,
        "ОБРАЗОВАТЕЛЬНОЕ УЧРЕЖДЕНИЕ ВЫСШЕГО ОБРАЗОВАНИЯ",
        bold=True,
        size=12,
        space_after=0,
    )
    add_center(doc, "«НАЦИОНАЛЬНЫЙ ИССЛЕДОВАТЕЛЬСКИЙ", bold=True, size=12, space_after=0)
    add_center(doc, "УНИВЕРСИТЕТ ИТМО»", bold=True, size=12, space_after=12)
    add_center(doc, "ФАКУЛЬТЕТ ПРОГРАММНОЙ ИНЖЕНЕРИИ", bold=True, size=12, space_after=0)
    add_center(doc, "И КОМПЬЮТЕРНОЙ ТЕХНИКИ", bold=True, size=12, space_after=36)

    for _ in range(2):
        doc.add_paragraph()

    add_center(
        doc,
        "Создание информационной системы для управления",
        bold=True,
        size=14,
        space_after=0,
    )
    add_center(doc, "работой автосервиса", bold=True, size=14, space_after=18)
    add_center(doc, "курсовая работа, этап №2", size=14, space_after=12)
    add_center(doc, "по дисциплине «Информационные системы»", size=14, space_after=0)
    add_center(
        doc,
        "направления бакалавриата 09.03.04 «Программная инженерия»",
        size=14,
        space_after=36,
    )

    for _ in range(4):
        doc.add_paragraph()

    add_center(doc, "г. Санкт-Петербург — 2026", size=14, space_after=0)
    doc.add_page_break()

    # ---------- Оглавление / Задание ----------
    add_heading_custom(doc, "Оглавление")
    for item in [
        "Задание",
        "1. ER-модель базы данных",
        "2. Даталогическая модель",
        "3. Реализация модели в PostgreSQL",
        "4. Обеспечение целостности данных (DDL и триггеры)",
        "5. Скрипты создания, удаления и заполнения БД",
        "6. PL/pgSQL-функции и процедуры",
        "7. Индексы и обоснование",
        "Вывод",
        "Приложение А. Фрагменты SQL-кода",
        "Приложение Б. Java-модель сущностей",
    ]:
        add_body(doc, item, first_line=False)

    add_heading_custom(doc, "Задание")
    for item in [
        "1. Сформировать ER-модель базы данных (на основе описаний предметной области и прецедентов из предыдущего этапа). ER-модель должна включать не менее 10 сущностей и содержать хотя бы одно отношение «многие-ко-многим».",
        "2. Согласовать ER-модель с преподавателем. На основе ER-модели построить даталогическую модель.",
        "3. Реализовать даталогическую модель в реляционной СУБД PostgreSQL.",
        "4. Обеспечить целостность данных при помощи средств языка DDL и триггеров.",
        "5. Реализовать скрипты для создания, удаления базы данных, заполнения базы тестовыми данными.",
        "6. Предложить PL/pgSQL-функции и процедуры для выполнения критически важных запросов.",
        "7. Создать индексы на основе анализа использования базы данных в контексте прецедентов этапа 1 и обосновать их полезность.",
        "8. Составить отчет.",
    ]:
        add_body(doc, item, first_line=False)

    # ---------- 1. ER ----------
    add_heading_custom(doc, "1. ER-модель базы данных")
    add_body(
        doc,
        "ER-модель построена по описанию предметной области и прецедентам этапа 1. "
        "Модель охватывает учёт пользователей и ролей, автомобилей, точек обслуживания, "
        "каталога услуг и запчастей, заявок, позиций запчастей, запросов на покупку, "
        "решений клиента, отчётов о работах и истории статусов.",
    )
    add_body(
        doc,
        "В модели 16 сущностей (таблиц): roles, users, vehicles, service_locations, "
        "mechanic_locations, service_types, parts, part_compatibilities, service_type_parts, "
        "request_statuses, service_requests, request_parts, part_purchase_requests, "
        "purchase_approvals, work_reports, request_status_history.",
    )
    add_body(
        doc,
        "Отношения «многие-ко-многим»:",
        first_line=True,
    )
    add_body(
        doc,
        "• mechanic_locations — механик может работать в нескольких точках, точка имеет нескольких механиков;",
        first_line=False,
    )
    add_body(
        doc,
        "• part_compatibilities — запчасть совместима с несколькими марками/моделями, модель может иметь много запчастей;",
        first_line=False,
    )
    add_body(
        doc,
        "• service_type_parts — услуга модификации связана с несколькими запчастями каталога.",
        first_line=False,
    )
    add_body(
        doc,
        "Код ER-модели представлен в файле model/er_model.dbml (нотация DBML) и в диаграмме "
        "diagrams/er_diagram.mmd. На рисунке 1 приведена схема связей сущностей.",
    )
    img = ROOT / "diagrams" / "er_diagram.png"
    if img.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run()
        run.add_picture(str(img), width=Cm(16))
    add_center(doc, "Рисунок 1 — ER-диаграмма информационной системы автосервиса", size=12)

    # ---------- 2. Datalogical ----------
    add_heading_custom(doc, "2. Даталогическая модель")
    add_body(
        doc,
        "Даталогическая модель получена нормализацией ER-модели до 3НФ. Каждая сущность "
        "отображена в отдельную реляционную таблицу схемы autoservice. Связи 1:N реализованы "
        "внешними ключами; связи M:N — отдельными связующими таблицами с составными первичными ключами.",
    )
    add_body(
        doc,
        "Ключевые преобразования: роль и статус заявки вынесены в справочники; режим согласования "
        "покупок хранится в заявке (ALLOW_WITHOUT_CONFIRMATION / REQUIRE_CONFIRMATION); "
        "аудит переходов статусов — в request_status_history; факт решения клиента — в purchase_approvals.",
    )
    add_body(
        doc,
        "Основные атрибуты service_requests: client_id, vehicle_id, service_type_id, location_id, "
        "mechanic_id, status_id, description, desired_date, agreed_date, purchase_approval_mode, "
        "cancel_reason, timestamps. Жизненный цикл статусов: NEW → CONFIRMED → ASSIGNED → "
        "IN_PROGRESS → COMPLETED; дополнительно AWAITING_PART_APPROVAL, AWAITING_PART, CANCELLED.",
    )
    add_body(
        doc,
        "Модель также записана кодом JPA-сущностей в каталоге model/java (для последующей "
        "реализации уровня хранения на Spring).",
    )

    # ---------- 3. PostgreSQL ----------
    add_heading_custom(doc, "3. Реализация модели в PostgreSQL")
    add_body(
        doc,
        "База данных развёрнута в PostgreSQL 16 (Docker-контейнер autoservice-db, порт хоста 5434). "
        "Создана схема autoservice, в ней созданы все 16 таблиц, ограничения PRIMARY KEY / "
        "FOREIGN KEY / UNIQUE / CHECK, индексы, триггеры и PL/pgSQL-объекты.",
    )
    add_body(
        doc,
        "Подключение: host=localhost, port=5434, db=autoservice, user=autoservice, password=autoservice. "
        "Инициализация выполняется скриптом sql/init_all.sql. Таблицы созданы и заполнены тестовыми данными; "
        "проверка выборками list_requests_filtered и get_vehicle_service_history подтвердила корректность сценариев.",
    )
    add_body(
        doc,
        "Фактическое состояние после инициализации: 4 заявки (COMPLETED, IN_PROGRESS, COMPLETED, NEW), "
        "6 пользователей, 2 активные точки, 4 автомобиля, каталог услуг и запчастей, отработанный "
        "сценарий покупки детали с подтверждением клиента.",
    )

    # ---------- 4. Integrity ----------
    add_heading_custom(doc, "4. Обеспечение целостности данных (DDL и триггеры)")
    add_body(
        doc,
        "На уровне DDL заданы внешние ключи, уникальность email/госномера/VIN/SKU, CHECK-ограничения "
        "на цены, количества, допустимые коды ролей, статусов, режимов согласования и источников запчастей.",
    )
    add_body(doc, "Триггеры (sql/02_triggers.sql) обеспечивают бизнес-целостность:", first_line=True)
    for t in [
        "• check_mechanic_role — в mechanic_locations можно добавить только пользователя с ролью MECHANIC;",
        "• check_vehicle_owner_is_client — владелец автомобиля только CLIENT;",
        "• check_service_type_is_modification — запчасти связываются только с услугами модификации;",
        "• validate_service_request — автомобиль принадлежит клиенту; точка активна; механик закреплён за точкой; допустимы только корректные переходы статусов; нельзя завершить заявку при открытых покупках;",
        "• log_request_status_change — аудит смены статусов;",
        "• validate_part_purchase_request / after_part_purchase_request — запрос создаёт механик заявки; при предварительном разрешении выполняется AUTO_APPROVED и переход в AWAITING_PART, иначе — ожидание решения клиента;",
        "• validate_purchase_approval — решение принимает только владелец заявки;",
        "• validate_work_report — отчёт пишет назначенный механик;",
        "• set_updated_at — автоматическое обновление метки времени.",
    ]:
        add_body(doc, t, first_line=False)

    # ---------- 5. Scripts ----------
    add_heading_custom(doc, "5. Скрипты создания, удаления и заполнения БД")
    add_body(doc, "Набор скриптов каталога sql/:", first_line=True)
    for s in [
        "• 00_drop.sql — удаление схемы autoservice CASCADE;",
        "• 01_create_schema.sql — создание таблиц и ограничений;",
        "• 02_triggers.sql — функции и триггеры целостности;",
        "• 03_functions.sql — PL/pgSQL API под прецеденты;",
        "• 04_indexes.sql — индексы;",
        "• 05_seed.sql — тестовые данные и демонстрационные сценарии;",
        "• init_all.sql — последовательный запуск всех скриптов.",
    ]:
        add_body(doc, s, first_line=False)
    add_body(
        doc,
        "Запуск: docker compose up -d && docker exec -i autoservice-db psql -U autoservice -d autoservice -f /sql/init_all.sql. "
        "Конфигурация контейнера — compose.yaml.",
    )

    # ---------- 6. Functions ----------
    add_heading_custom(doc, "6. PL/pgSQL-функции и процедуры")
    add_body(
        doc,
        "Реализованы функции и процедуры, необходимые для прецедентов этапа 1 и последующей реализации "
        "уровня хранения (этап 3):",
    )
    for f in [
        "• register_client — регистрация клиента (C000);",
        "• add_vehicle — добавление автомобиля (C002);",
        "• create_service_request — создание заявки с проверкой совместимости запчасти для модификации (C003, N014);",
        "• confirm_and_assign_request — подтверждение и назначение механика (M001/M002);",
        "• start_work — перевод ASSIGNED → IN_PROGRESS (K002);",
        "• create_part_purchase_request — запрос на покупку детали (K005);",
        "• decide_purchase — решение клиента (C011);",
        "• mark_part_purchased_and_received — покупка/получение детали и возврат в работу (K006/K007);",
        "• complete_request — отчёт и завершение заявки (K003/K004);",
        "• get_vehicle_service_history — история обслуживания автомобиля (C006/C012);",
        "• find_compatible_parts — подбор совместимых запчастей (C008);",
        "• list_requests_filtered — фильтрованный список заявок для менеджера (M000).",
    ]:
        add_body(doc, f, first_line=False)

    # ---------- 7. Indexes ----------
    add_heading_custom(doc, "7. Индексы и обоснование")
    add_body(
        doc,
        "Индексы созданы исходя из типовых запросов прецедентов этапа 1. Ниже — соответствие "
        "индекса бизнес-процессу.",
    )
    for idx in [
        "• idx_users_email_lower — вход по email (C001);",
        "• idx_vehicles_owner_id, idx_service_requests_vehicle_id — автомобили клиента и история по ТС (C002, C006);",
        "• idx_service_requests_client_id — список заявок клиента (C004);",
        "• idx_service_requests_status_id / desired_date / mechanic_id / location_id — фильтры менеджера (M000);",
        "• idx_service_requests_mechanic_status — очередь работ механика (K000);",
        "• idx_part_compat_make_model, idx_part_compat_part_id — поиск совместимых деталей (C008, N014);",
        "• idx_mechanic_locations_location — проверка права механика работать в точке (M002, N015);",
        "• idx_ppr_request_status, idx_ppr_status_pending — контроль открытых закупок (M008, C011, N008);",
        "• idx_request_status_history_request_id, idx_request_parts_request_id — аудит и детализация заявки (M005, C012).",
    ]:
        add_body(doc, idx, first_line=False)
    add_body(
        doc,
        "Частичные индексы по активным запчастям и открытым закупкам уменьшают объём индекса "
        "и ускоряют горячие выборки без ущерба для редких архивных записей.",
    )

    # ---------- Вывод ----------
    add_heading_custom(doc, "Вывод")
    add_body(
        doc,
        "В ходе второго этапа спроектирована и реализована база данных ИС сети автосервиса. "
        "ER-модель содержит 16 сущностей и три связи «многие-ко-многим», отражает каталог запчастей, "
        "точки обслуживания, режимы согласования покупок и полный жизненный цикл заявки. "
        "Даталогическая модель воплощена в PostgreSQL: созданы таблицы, ограничения DDL, триггеры "
        "целостности, индексы под прецеденты этапа 1, PL/pgSQL-функции/процедуры и скрипты "
        "инициализации/удаления/заполнения. База развёрнута локально в Docker и готова к использованию "
        "на этапе 3 при реализации уровня хранения на Spring.",
    )

    # ---------- Appendix A ----------
    doc.add_page_break()
    add_heading_custom(doc, "Приложение А. Фрагменты SQL-кода")
    add_body(doc, "А.1 Создание ключевых таблиц (фрагмент 01_create_schema.sql)", first_line=False)
    add_code(
        doc,
        """CREATE TABLE service_requests (
    id                      BIGSERIAL PRIMARY KEY,
    client_id               BIGINT NOT NULL REFERENCES users(id),
    vehicle_id              BIGINT NOT NULL REFERENCES vehicles(id),
    service_type_id         BIGINT NOT NULL REFERENCES service_types(id),
    location_id             BIGINT NOT NULL REFERENCES service_locations(id),
    mechanic_id             BIGINT REFERENCES users(id),
    status_id               SMALLINT NOT NULL REFERENCES request_statuses(id),
    description             TEXT NOT NULL,
    desired_date            DATE NOT NULL,
    agreed_date             DATE,
    purchase_approval_mode  VARCHAR(40) NOT NULL,
    ...
);""",
    )
    add_body(doc, "А.2 Пример PL/pgSQL-процедуры confirm_and_assign_request", first_line=False)
    add_code(
        doc,
        """CREATE OR REPLACE PROCEDURE confirm_and_assign_request(
    p_request_id BIGINT,
    p_mechanic_id BIGINT,
    p_agreed_date DATE
) LANGUAGE plpgsql AS $$
BEGIN
    UPDATE service_requests
    SET status_id = (SELECT id FROM request_statuses WHERE code = 'CONFIRMED'),
        agreed_date = p_agreed_date
    WHERE id = p_request_id;

    UPDATE service_requests
    SET mechanic_id = p_mechanic_id,
        status_id = (SELECT id FROM request_statuses WHERE code = 'ASSIGNED')
    WHERE id = p_request_id;
END;
$$;""",
    )
    add_body(
        doc,
        "Полные исходные тексты приведены в каталоге sql/ проекта "
        "/home/vadim/is-coursework.",
        first_line=False,
    )

    # ---------- Appendix B ----------
    add_heading_custom(doc, "Приложение Б. Java-модель сущностей")
    add_body(
        doc,
        "Для согласования с архитектурой этапа 1 (Spring MVC + JPA) сущности описаны Java-классами "
        "в model/java. Пример фрагмента ServiceRequest:",
        first_line=False,
    )
    add_code(
        doc,
        """@Entity
@Table(name = "service_requests", schema = "autoservice")
public class ServiceRequest {
    @Id @GeneratedValue(strategy = GenerationType.IDENTITY)
    private Long id;
    @ManyToOne @JoinColumn(name = "client_id") private User client;
    @ManyToOne @JoinColumn(name = "vehicle_id") private Vehicle vehicle;
    @ManyToOne @JoinColumn(name = "location_id") private ServiceLocation location;
    @ManyToOne @JoinColumn(name = "mechanic_id") private User mechanic;
    @ManyToOne @JoinColumn(name = "status_id") private RequestStatus status;
    private String purchaseApprovalMode;
}""",
    )

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(f"Saved: {OUT}")


if __name__ == "__main__":
    build()
